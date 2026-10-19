"""Build the SupplierPulse Excel workbook.

Design rules
------------
* Navy column headers = data loaded from the pipeline (values).
  Teal column headers  = live Excel formulas.
  The Instructions sheet states this, so values are never presented as formulas.
* All synthetic delivery metrics, the scorecard, exception flags and the scenario model are
  recalculated by Excel from the line-level table `tblLines` (COUNTIFS / SUMIFS / AVERAGEIFS /
  XLOOKUP / INDEX-MATCH), using thresholds on the Assumptions sheet.
* Observed City of Austin sheets hold aggregated values from the pipeline with formula-driven
  shares, ranks, Pareto classes and flags (the 321k-line source is not embedded).
* Formulas use plain absolute ranges (not structured references) so they also evaluate in
  LibreOffice and in the Python `formulas` engine used by the test suite.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "1F2A44"
SLATE = "475569"
TEAL = "0F766E"
LIGHT = "F1F5F9"
AMBER_FILL = "FEF3C7"
RED_FILL = "FEE2E2"
GREEN_FILL = "DCFCE7"
INPUT_FILL = "FFF7E6"

HDR_DATA = PatternFill("solid", fgColor=NAVY)
HDR_FORMULA = PatternFill("solid", fgColor=TEAL)
WHITE_BOLD = Font(color="FFFFFF", bold=True)
TITLE = Font(size=16, bold=True, color=NAVY)
SUBTITLE = Font(size=10, italic=True, color=SLATE)
BOLD = Font(bold=True, color=NAVY)
THIN = Side(style="thin", color="CBD5E1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PCT = "0.0%"
USD = '"$"#,##0'
USD2 = '"$"#,##0.00'
NUM = "#,##0"
NUM1 = "#,##0.0"
DATE = "yyyy-mm-dd"

SYN_LABEL = "SYNTHETIC DEMONSTRATION DATA - not real suppliers or real performance"
OBS_LABEL = "OBSERVED DATA - City of Austin purchase orders (data.austintexas.gov), {period}"


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE
    ws.sheet_view.showGridLines = False


def _header(ws, row: int, col: int, text: str, formula: bool = False) -> None:
    c = ws.cell(row=row, column=col, value=text)
    c.fill = HDR_FORMULA if formula else HDR_DATA
    c.font = WHITE_BOLD
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BOX


def _widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


class Book:
    """Small helper that tracks named input cells and table ranges."""

    def __init__(self):
        self.wb = Workbook()
        self.cells: dict[str, str] = {}
        self.lines_cols: dict[str, str] = {}
        self.lines_last = 0

    def L(self, key: str) -> str:
        col = self.lines_cols[key]
        return f"Syn_Lines!${col}$2:${col}${self.lines_last}"


# ---------------------------------------------------------------------------------------------

def build_workbook(r, path: Path | None = None) -> Path:
    s = r.settings
    path = path or s.path("deliverables_dir") / "SupplierPulse_Procurement_Analysis.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    b = Book()
    wb = b.wb
    period = r.headline["period_label"]

    ws_intro = wb.active
    ws_intro.title = "Instructions"
    _assumptions(b, r)
    _lines(b, r)
    _suppliers(b, r)
    _summary(b, r)
    _dependency(b, r)
    _scorecard(b, r)
    _trend(b, r)
    _exceptions(b, r)
    _scenario(b, r)
    _austin_vendors(b, r, period)
    _austin_categories(b, r, period)
    _austin_prices(b, r, period)
    _actions(b, r)
    _dictionary(b, r)
    _validation(b, r)
    _instructions(ws_intro, b, r, period)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------------------------

def _assumptions(b: Book, r) -> None:
    s = r.settings
    ws = b.wb.create_sheet("Assumptions")
    _title(ws, "Assumptions and thresholds", "Yellow cells are inputs. Every formula in the workbook reads its thresholds from this sheet.")
    sc, ex, scen, syn = s.scorecard, s.exceptions, s.scenario, s.synthetic
    rows = [
        ("section", "Reporting (synthetic delivery data)"),
        ("as_of", "As-of (reporting) date", pd.Timestamp(syn["as_of_date"]).to_pydatetime(), DATE, "Lines promised after this date are not yet due"),
        ("grace", "On-time grace days", syn["on_time_grace_days"], NUM, "Days after promised date still counted on time"),
        ("tol", "In-full tolerance", syn["in_full_tolerance"], PCT, "0% = full ordered quantity required"),
        ("section", "Scorecard weights (must total 100%)"),
        ("w_otif", "Weight: OTIF", sc["weight_otif"], PCT, ""),
        ("w_quality", "Weight: quality (reject rate)", sc["weight_quality"], PCT, ""),
        ("w_lt", "Weight: lead-time consistency", sc["weight_leadtime_consistency"], PCT, ""),
        ("w_late", "Weight: lateness (avg days late)", sc["weight_lateness"], PCT, ""),
        ("section", "Scorecard scaling (target = 100 points, poor = 0 points)"),
        ("otif_t", "OTIF target", sc["otif_target"], PCT, ""),
        ("otif_p", "OTIF poor", sc["otif_poor"], PCT, ""),
        ("rej_t", "Reject rate target", sc["reject_rate_target"], PCT, ""),
        ("rej_p", "Reject rate poor", sc["reject_rate_poor"], PCT, ""),
        ("lt_t", "Lead-time SD target (days)", sc["leadtime_sd_target"], NUM1, ""),
        ("lt_p", "Lead-time SD poor (days)", sc["leadtime_sd_poor"], NUM1, ""),
        ("late_t", "Avg days late target", sc["avg_days_late_target"], NUM1, ""),
        ("late_p", "Avg days late poor", sc["avg_days_late_poor"], NUM1, ""),
        ("section", "Evidence and tiers"),
        ("min_ranked", "Due lines needed to rank", sc["min_lines_ranked"], NUM, ""),
        ("min_prov", "Due lines needed for a provisional score", sc["min_lines_provisional"], NUM, ""),
        ("min_completed", "Completed lines needed for lead-time variability", ex["min_completed_for_variability"], NUM, ""),
        ("tier_pref", "Tier: Preferred at or above", sc["tier_preferred"], NUM, ""),
        ("tier_acc", "Tier: Acceptable at or above", sc["tier_acceptable"], NUM, ""),
        ("tier_imp", "Tier: Improvement needed at or above", sc["tier_improve"], NUM, ""),
        ("section", "Exception thresholds"),
        ("x_late", "Chronic late: OTIF below", ex["chronic_late_otif"], PCT, "S1"),
        ("x_min", "Minimum due lines for rate-based flags", ex["min_lines_for_rate_flags"], NUM, "S1, S2, S4, S5"),
        ("x_lt", "Unstable lead time: SD at or above (days)", ex["unstable_leadtime_sd"], NUM1, "S3"),
        ("x_if", "Incomplete deliveries: in-full below", ex["incomplete_in_full_rate"], PCT, "S4"),
        ("x_rej", "Quality: reject rate at or above", ex["quality_reject_rate"], PCT, "S5"),
        ("x_det", "Deterioration: OTIF drop of at least", ex["deterioration_points"], PCT, "S2"),
        ("x_ppv", "PPV: paid above contract by at least", ex["ppv_flag_ratio"], PCT, "S6"),
        ("dep_hc", "High dependency: category share at or above", ex["high_dependency_category_share"], PCT, "S8"),
        ("dep_ht", "High dependency: share of total at or above", ex["high_dependency_total_share"], PCT, "S8"),
        ("dep_mc", "Medium dependency: category share at or above", ex["medium_dependency_category_share"], PCT, ""),
        ("dep_mt", "Medium dependency: share of total at or above", ex["medium_dependency_total_share"], PCT, ""),
        ("section", "Scenario assumptions (modelled, not realised)"),
        ("cost_late", "Cost of handling one late PO line (USD)", scen["cost_per_late_line"], USD, "Expediting and disruption; an assumption"),
    ]
    row = 4
    _header(ws, row, 2, "Assumption")
    _header(ws, row, 3, "Value")
    _header(ws, row, 4, "Used by / note")
    row += 1
    for item in rows:
        if item[0] == "section":
            ws.cell(row=row, column=2, value=item[1]).font = BOLD
            row += 1
            continue
        key, label, value, fmt, note = item
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.number_format = fmt
        c.fill = PatternFill("solid", fgColor=INPUT_FILL)
        c.border = BOX
        ws.cell(row=row, column=4, value=note).font = Font(color=SLATE)
        b.cells[key] = f"Assumptions!$C${row}"
        row += 1
    ws.cell(row=row + 1, column=2, value="Weights total (check)").font = BOLD
    c = ws.cell(row=row + 1, column=3, value=f"={b.cells['w_otif']}+{b.cells['w_quality']}+{b.cells['w_lt']}+{b.cells['w_late']}")
    c.number_format = PCT
    b.cells["w_total"] = f"Assumptions!$C${row + 1}"
    _widths(ws, {"A": 2, "B": 48, "C": 16, "D": 44})


def _lines(b: Book, r) -> None:
    ws = b.wb.create_sheet("Syn_Lines")
    f = r.facts.sort_values("po_line_id").reset_index(drop=True)
    last_receipt = f["last_receipt_date"].fillna(f["short_close_date"])
    data_cols = [
        ("po_line_id", f["po_line_id"], None), ("po_id", f["po_id"], None), ("supplier_id", f["supplier_id"], None),
        ("category_id", f["category_id"], None), ("item_id", f["item_id"], None),
        ("order_date", f["order_date"], DATE), ("promised_date", f["promised_date"], DATE),
        ("ordered_qty", f["ordered_qty"], NUM), ("unit_price", f["unit_price"], USD2), ("standard_price", f["standard_price"], USD2),
        ("cancelled", f["cancelled"].astype(int), "0"), ("short_closed", f["short_closed"].astype(int), "0"),
        ("received_qty", f["received_qty"], NUM), ("qty_by_promise", f["qty_by_promise"], NUM), ("rejected_qty", f["rejected_qty"], NUM),
        ("first_receipt_date", f["first_receipt_date"], DATE), ("completion_date", f["completion_date"], DATE),
        ("last_receipt_or_close_date", last_receipt, DATE),
    ]
    for i, (name, _, _) in enumerate(data_cols, start=1):
        b.lines_cols[name] = get_column_letter(i)
    formula_names = ["line_value", "standard_value", "due", "on_time", "in_full", "otif", "complete", "open", "overdue",
                     "days_late", "lead_time_days", "lead_time_sq", "open_value", "due_period", "due_month"]
    for j, name in enumerate(formula_names, start=len(data_cols) + 1):
        b.lines_cols[name] = get_column_letter(j)
    C = b.lines_cols
    a = b.cells
    n = len(f)
    b.lines_last = n + 1

    for i, (name, _, _) in enumerate(data_cols, start=1):
        _header(ws, 1, i, name)
    for j, name in enumerate(formula_names, start=len(data_cols) + 1):
        _header(ws, 1, j, name, formula=True)

    cols_values = [series.tolist() for _, series, _ in data_cols]
    for i in range(n):
        row = i + 2
        for k, (name, _, fmt) in enumerate(data_cols):
            v = cols_values[k][i]
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
                v = None
            elif hasattr(v, "item"):
                v = v.item()
            ws.cell(row=row, column=k + 1, value=v)
        g = lambda key: f"{C[key]}{row}"  # noqa: E731
        formulas = {
            "line_value": f"={g('ordered_qty')}*{g('unit_price')}",
            "standard_value": f"={g('ordered_qty')}*{g('standard_price')}",
            "due": f"=IF(AND({g('cancelled')}=0,{g('promised_date')}<={a['as_of']}),1,0)",
            "on_time": f"=IF({g('due')}=1,IF({g('received_qty')}>0,IF({g('first_receipt_date')}<={g('promised_date')}+{a['grace']},1,0),0),0)",
            "in_full": f"=IF(AND({g('due')}=1,{g('received_qty')}>={g('ordered_qty')}*(1-{a['tol']})),1,0)",
            "otif": f"=IF(AND({g('due')}=1,{g('qty_by_promise')}>={g('ordered_qty')}*(1-{a['tol']})),1,0)",
            "complete": f"=IF(ISNUMBER({g('completion_date')}),1,0)",
            "open": f"=IF(AND({g('cancelled')}=0,{g('short_closed')}=0,{g('complete')}=0),1,0)",
            "overdue": f"=IF(AND({g('open')}=1,{g('promised_date')}<{a['as_of']}),1,0)",
            "days_late": (f"=IF({g('due')}=0,\"\",IF({g('complete')}=1,MAX(0,{g('completion_date')}-{g('promised_date')}),"
                          f"IF({g('short_closed')}=1,MAX(0,{g('last_receipt_or_close_date')}-{g('promised_date')}),"
                          f"MAX(0,{a['as_of']}-{g('promised_date')}))))"),
            "lead_time_days": f"=IF({g('complete')}=1,{g('completion_date')}-{g('order_date')},\"\")",
            "lead_time_sq": f"=IF({g('complete')}=1,{g('lead_time_days')}^2,\"\")",
            "open_value": f"=IF({g('open')}=1,({g('ordered_qty')}-{g('received_qty')})*{g('unit_price')},0)",
            "due_period": f"=IF({g('promised_date')}<EDATE({a['as_of']},-12)+1,\"prior_12m\",\"last_12m\")",
            "due_month": f"=DATE(YEAR({g('promised_date')}),MONTH({g('promised_date')}),1)",
        }
        for name in formula_names:
            ws[f"{C[name]}{row}"] = formulas[name]

    fmt_map = {name: fmt for name, _, fmt in data_cols}
    fmt_map.update({"line_value": USD2, "standard_value": USD2, "open_value": USD2, "due_month": "yyyy-mm"})
    for name, fmt in fmt_map.items():
        if fmt:
            for row in range(2, n + 2):
                ws[f"{C[name]}{row}"].number_format = fmt
    last_col = get_column_letter(len(data_cols) + len(formula_names))
    tab = Table(displayName="tblLines", ref=f"A1:{last_col}{n + 1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "B2"
    for i in range(1, len(data_cols) + len(formula_names) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14


def _suppliers(b: Book, r) -> None:
    ws = b.wb.create_sheet("Syn_Suppliers")
    _title(ws, "Synthetic supplier master", SYN_LABEL)
    heads = ["supplier_id", "supplier_name", "primary_category_id", "categories_served", "design_note"]
    for j, h in enumerate(heads, start=1):
        _header(ws, 4, j, h)
    for i, rec in enumerate(r.syn.suppliers[heads].itertuples(index=False), start=5):
        for j, v in enumerate(rec, start=1):
            ws.cell(row=i, column=j, value=v)
    b.cells["sup_first"], b.cells["sup_last"] = 5, 4 + len(r.syn.suppliers)
    _widths(ws, {"A": 12, "B": 26, "C": 18, "D": 18, "E": 80})


SUMMARY_COLS = [
    # key, header, number format, formula template ({sup} = supplier cell, {r} = row)
    ("due_lines", "Due lines", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_due},1)"),
    ("otif_lines", "OTIF lines", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_otif},1)"),
    ("otif_rate", "OTIF rate", PCT, "=IF({due_lines}>0,{otif_lines}/{due_lines},\"\")"),
    ("on_time_rate", "On-time rate", PCT, "=IF({due_lines}>0,COUNTIFS({L_supplier_id},{sup},{L_on_time},1)/{due_lines},\"\")"),
    ("in_full_rate", "In-full rate", PCT, "=IF({due_lines}>0,COUNTIFS({L_supplier_id},{sup},{L_in_full},1)/{due_lines},\"\")"),
    ("late_lines", "Late lines", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_due},1,{L_days_late},\">0\")"),
    ("avg_days_late", "Avg days late (all due lines)", NUM1, "=IF({due_lines}>0,AVERAGEIFS({L_days_late},{L_supplier_id},{sup},{L_due},1),\"\")"),
    ("completed", "Completed lines", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_complete},1)"),
    ("avg_lt", "Avg lead time (days)", NUM1, "=IF({completed}>0,AVERAGEIFS({L_lead_time_days},{L_supplier_id},{sup},{L_complete},1),\"\")"),
    ("lt_sd", "Lead-time SD (days)", NUM1,
     "=IF({completed}>1,SQRT(MAX(0,(SUMIFS({L_lead_time_sq},{L_supplier_id},{sup},{L_complete},1)-{completed}*{avg_lt}^2)/({completed}-1))),\"\")"),
    ("received", "Received qty", NUM, "=SUMIFS({L_received_qty},{L_supplier_id},{sup},{L_cancelled},0)"),
    ("rejected", "Rejected qty", NUM, "=SUMIFS({L_rejected_qty},{L_supplier_id},{sup},{L_cancelled},0)"),
    ("reject_rate", "Reject rate", "0.00%", "=IF({received}>0,{rejected}/{received},\"\")"),
    ("spend", "Spend (ordered value)", USD, "=SUMIFS({L_line_value},{L_supplier_id},{sup},{L_cancelled},0)"),
    ("std_value", "Value at contract price", USD, "=SUMIFS({L_standard_value},{L_supplier_id},{sup},{L_cancelled},0)"),
    ("ppv_pct", "PPV vs contract", PCT, "=IF({std_value}>0,{spend}/{std_value}-1,\"\")"),
    ("share_total", "Share of total spend", PCT, "={spend}/SUM({spend_col})"),
    ("open_lines", "Open lines", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_open},1)"),
    ("overdue_lines", "Overdue lines", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_overdue},1)"),
    ("overdue_value", "Overdue open value", USD, "=SUMIFS({L_open_value},{L_supplier_id},{sup},{L_overdue},1)"),
    ("due_prior", "Due lines prior 12m", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_due},1,{L_due_period},\"prior_12m\")"),
    ("otif_prior", "OTIF prior 12m", PCT,
     "=IF({due_prior}>0,COUNTIFS({L_supplier_id},{sup},{L_otif},1,{L_due_period},\"prior_12m\")/{due_prior},\"\")"),
    ("due_last", "Due lines last 12m", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_due},1,{L_due_period},\"last_12m\")"),
    ("otif_last", "OTIF last 12m", PCT,
     "=IF({due_last}>0,COUNTIFS({L_supplier_id},{sup},{L_otif},1,{L_due_period},\"last_12m\")/{due_last},\"\")"),
    ("otif_change", "OTIF change (pts)", PCT, "=IF(OR({otif_prior}=\"\",{otif_last}=\"\"),\"\",{otif_last}-{otif_prior})"),
    ("short_closed", "Lines closed short", NUM, "=COUNTIFS({L_supplier_id},{sup},{L_short_closed},1)"),
]


def _summary(b: Book, r) -> None:
    ws = b.wb.create_sheet("Supplier_Summary")
    _title(ws, "Supplier delivery, quality and spend summary", SYN_LABEL + ". Every metric is a live formula over Syn_Lines.")
    ws["A3"] = "Measured per PO line. Only due lines (not cancelled, promised on/before the as-of date) enter delivery rates; unreceived due lines count as late."
    ws["A3"].font = Font(color=SLATE, size=9)
    _header(ws, 4, 1, "Supplier ID")
    col_of = {}
    for j, (key, head, fmt, _) in enumerate(SUMMARY_COLS, start=2):
        _header(ws, 4, j, head, formula=True)
        col_of[key] = get_column_letter(j)
    sups = r.syn.suppliers["supplier_id"].tolist()
    first, last = 5, 4 + len(sups)
    b.cells["sum_first"], b.cells["sum_last"] = first, last
    b.cells["sum_cols"] = col_of
    lrefs = {f"L_{k}": b.L(k) for k in b.lines_cols}
    for i, sup in enumerate(sups):
        row = first + i
        ws.cell(row=row, column=1, value=sup).font = BOLD
        refs = {k: f"{c}{row}" for k, c in col_of.items()}
        for j, (key, _, fmt, tmpl) in enumerate(SUMMARY_COLS, start=2):
            cell = ws.cell(row=row, column=j, value=tmpl.format(sup=f"$A{row}", spend_col=f"${col_of['spend']}${first}:${col_of['spend']}${last}", **lrefs, **refs))
            cell.number_format = fmt
    tot = last + 1
    ws.cell(row=tot, column=1, value="All suppliers").font = BOLD
    for key in ["due_lines", "otif_lines", "late_lines", "completed", "received", "rejected", "spend", "std_value", "open_lines", "overdue_lines", "overdue_value", "short_closed"]:
        c = col_of[key]
        ws[f"{c}{tot}"] = f"=SUM({c}{first}:{c}{last})"
        ws[f"{c}{tot}"].number_format = dict((k, f) for k, _, f, _ in SUMMARY_COLS)[key]
        ws[f"{c}{tot}"].font = BOLD
    ws[f"{col_of['otif_rate']}{tot}"] = f"={col_of['otif_lines']}{tot}/{col_of['due_lines']}{tot}"
    ws[f"{col_of['otif_rate']}{tot}"].number_format = PCT
    ws[f"{col_of['reject_rate']}{tot}"] = f"={col_of['rejected']}{tot}/{col_of['received']}{tot}"
    ws[f"{col_of['reject_rate']}{tot}"].number_format = "0.00%"
    b.cells["sum_total_row"] = tot
    rng = f"{col_of['otif_rate']}{first}:{col_of['otif_rate']}{last}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0.6, start_color="F8D7DA",
                                                      mid_type="num", mid_value=0.85, mid_color="FFFFFF",
                                                      end_type="num", end_value=1.0, end_color="D1F2EB"))
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 12
    for j in range(2, len(SUMMARY_COLS) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.row_dimensions[4].height = 45


def _dependency(b: Book, r) -> None:
    ws = b.wb.create_sheet("Dependency")
    _title(ws, "Supplier dependency by category (spend share)", SYN_LABEL)
    cats = r.syn.categories["category_id"].tolist()
    sups = r.syn.suppliers["supplier_id"].tolist()
    _header(ws, 4, 1, "Supplier ID")
    for j, cat in enumerate(cats, start=2):
        _header(ws, 4, j, f"Spend: {cat}", formula=True)
    share_start = len(cats) + 3
    for j, cat in enumerate(cats, start=share_start):
        _header(ws, 4, j, f"Share: {cat}", formula=True)
    extra = share_start + len(cats)
    heads = ["Max category share", "Main category", "Suppliers in main category", "Share of total spend", "Dependency risk"]
    for k, h in enumerate(heads):
        _header(ws, 4, extra + k, h, formula=True)
    first, last = 5, 4 + len(sups)
    count_row = last + 2
    total_row = last + 1
    a = b.cells
    for i, sup in enumerate(sups):
        row = first + i
        ws.cell(row=row, column=1, value=sup).font = BOLD
        for j, cat in enumerate(cats, start=2):
            c = ws.cell(row=row, column=j, value=f"=SUMIFS({b.L('line_value')},{b.L('supplier_id')},$A{row},{b.L('category_id')},\"{cat}\",{b.L('cancelled')},0)")
            c.number_format = USD
        for k, cat in enumerate(cats):
            spend_col = get_column_letter(2 + k)
            c = ws.cell(row=row, column=share_start + k, value=f"=IF({spend_col}${total_row}>0,{spend_col}{row}/{spend_col}${total_row},0)")
            c.number_format = PCT
        s0, s1 = get_column_letter(share_start), get_column_letter(share_start + len(cats) - 1)
        mx = get_column_letter(extra)
        ws.cell(row=row, column=extra, value=f"=MAX({s0}{row}:{s1}{row})").number_format = PCT
        # Main category = header of the largest share ("Share: CAT" -> "CAT" via MID); INDEX/MATCH lookup
        ws.cell(row=row, column=extra + 1, value=f"=MID(INDEX(${s0}$4:${s1}$4,MATCH({mx}{row},{s0}{row}:{s1}{row},0)),8,10)")
        ws.cell(row=row, column=extra + 2, value=f"=INDEX(${get_column_letter(2)}${count_row}:${get_column_letter(1 + len(cats))}${count_row},MATCH({mx}{row},{s0}{row}:{s1}{row},0))")
        sumc = b.cells["sum_cols"]["share_total"]
        ws.cell(row=row, column=extra + 3, value=f"=Supplier_Summary!{sumc}{b.cells['sum_first'] + i}").number_format = PCT
        dep = (f"=IF(OR({get_column_letter(extra + 2)}{row}=1,{mx}{row}>={a['dep_hc']},{get_column_letter(extra + 3)}{row}>={a['dep_ht']}),\"High\","
               f"IF(OR({mx}{row}>={a['dep_mc']},{get_column_letter(extra + 3)}{row}>={a['dep_mt']}),\"Medium\",\"Low\"))")
        ws.cell(row=row, column=extra + 4, value=dep)
    ws.cell(row=total_row, column=1, value="Category total").font = BOLD
    ws.cell(row=count_row, column=1, value="Suppliers with spend").font = BOLD
    for j in range(2, 2 + len(cats)):
        col = get_column_letter(j)
        ws[f"{col}{total_row}"] = f"=SUM({col}{first}:{col}{last})"
        ws[f"{col}{total_row}"].number_format = USD
        ws[f"{col}{count_row}"] = f"=COUNTIF({col}{first}:{col}{last},\">0\")"
    b.cells["dep_cols"] = {"max_share": get_column_letter(extra), "main_cat": get_column_letter(extra + 1),
                           "sup_in_cat": get_column_letter(extra + 2), "share_total": get_column_letter(extra + 3),
                           "dependency": get_column_letter(extra + 4)}
    b.cells["dep_first"] = first
    rng = f"{get_column_letter(share_start)}{first}:{get_column_letter(share_start + len(cats) - 1)}{last}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                      end_type="num", end_value=1, end_color="9FB4CC"))
    ws.column_dimensions["A"].width = 20
    for j in range(2, extra + 5):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.row_dimensions[4].height = 40


def _scorecard(b: Book, r) -> None:
    ws = b.wb.create_sheet("Scorecard")
    _title(ws, "Supplier performance scorecard", SYN_LABEL + ". Score = weighted 0-100 components; see Assumptions.")
    a = b.cells
    sc = a["sum_cols"]
    dc = a["dep_cols"]
    heads = [("Supplier ID", False), ("Supplier name (XLOOKUP)", True), ("Due lines", True), ("Evidence status", True),
             ("OTIF rate", True), ("Reject rate", True), ("Lead-time SD used", True), ("Avg days late", True),
             ("Score: OTIF", True), ("Score: quality", True), ("Score: lead-time consistency", True), ("Score: lateness", True),
             ("Weight used", True), ("Performance score", True), ("Rank (ranked only)", True), ("Performance tier", True),
             ("Dependency risk", True), ("Action quadrant", True)]
    for j, (h, fml) in enumerate(heads, start=1):
        _header(ws, 4, j, h, formula=fml)
    sups = r.syn.suppliers["supplier_id"].tolist()
    first, last = 5, 4 + len(sups)
    b.cells["sc_first"], b.cells["sc_last"] = first, last
    sf, sl = a["sup_first"], a["sup_last"]
    for i, sup in enumerate(sups):
        row = first + i
        srow = a["sum_first"] + i
        drow = a["dep_first"] + i
        S = lambda k: f"Supplier_Summary!{sc[k]}{srow}"  # noqa: E731
        ws.cell(row=row, column=1, value=sup).font = BOLD
        ws.cell(row=row, column=2, value=f"=XLOOKUP(A{row},Syn_Suppliers!$A${sf}:$A${sl},Syn_Suppliers!$B${sf}:$B${sl},\"not found\")")
        ws.cell(row=row, column=3, value=f"={S('due_lines')}").number_format = NUM
        ws.cell(row=row, column=4, value=f"=IF(C{row}>={a['min_ranked']},\"Ranked\",IF(C{row}>={a['min_prov']},\"Provisional\",\"Insufficient evidence\"))")
        ws.cell(row=row, column=5, value=f"={S('otif_rate')}").number_format = PCT
        ws.cell(row=row, column=6, value=f"={S('reject_rate')}").number_format = "0.00%"
        ws.cell(row=row, column=7, value=f"=IF({S('completed')}>={a['min_completed']},{S('lt_sd')},\"\")").number_format = NUM1
        ws.cell(row=row, column=8, value=f"={S('avg_days_late')}").number_format = NUM1
        ws.cell(row=row, column=9, value=f"=IF(E{row}=\"\",\"\",MIN(1,MAX(0,(E{row}-{a['otif_p']})/({a['otif_t']}-{a['otif_p']})))*100)")
        ws.cell(row=row, column=10, value=f"=IF(F{row}=\"\",\"\",MIN(1,MAX(0,({a['rej_p']}-F{row})/({a['rej_p']}-{a['rej_t']})))*100)")
        ws.cell(row=row, column=11, value=f"=IF(G{row}=\"\",\"\",MIN(1,MAX(0,({a['lt_p']}-G{row})/({a['lt_p']}-{a['lt_t']})))*100)")
        ws.cell(row=row, column=12, value=f"=IF(H{row}=\"\",\"\",MIN(1,MAX(0,({a['late_p']}-H{row})/({a['late_p']}-{a['late_t']})))*100)")
        ws.cell(row=row, column=13, value=(f"=IF(I{row}=\"\",0,{a['w_otif']})+IF(J{row}=\"\",0,{a['w_quality']})"
                                           f"+IF(K{row}=\"\",0,{a['w_lt']})+IF(L{row}=\"\",0,{a['w_late']})")).number_format = PCT
        ws.cell(row=row, column=14, value=(f"=IF(OR(D{row}=\"Insufficient evidence\",M{row}=0),\"\",ROUND((IF(I{row}=\"\",0,I{row}*{a['w_otif']})"
                                           f"+IF(J{row}=\"\",0,J{row}*{a['w_quality']})+IF(K{row}=\"\",0,K{row}*{a['w_lt']})"
                                           f"+IF(L{row}=\"\",0,L{row}*{a['w_late']}))/M{row},1))")).number_format = "0.0"
        ws.cell(row=row, column=15, value=f"=IF(D{row}=\"Ranked\",1+COUNTIFS($D${first}:$D${last},\"Ranked\",$N${first}:$N${last},\">\"&N{row}),\"\")")
        ws.cell(row=row, column=16, value=(f"=IF(N{row}=\"\",\"Not scored\",IF(N{row}>={a['tier_pref']},\"Preferred\",IF(N{row}>={a['tier_acc']},"
                                           f"\"Acceptable\",IF(N{row}>={a['tier_imp']},\"Improvement needed\",\"At risk\"))))"))
        ws.cell(row=row, column=17, value=f"=Dependency!{dc['dependency']}{drow}")
        ws.cell(row=row, column=18, value=(f"=IF(D{row}=\"Insufficient evidence\",\"Insufficient evidence\",IF(Q{row}=\"High\","
                                           f"IF(N{row}>={a['tier_acc']},\"Strategic - protect and monitor\",\"Critical - act now\"),"
                                           f"IF(N{row}>={a['tier_acc']},\"Maintain\",\"Develop or replace\")))"))
        for col in "IJKL":
            ws[f"{col}{row}"].number_format = "0.0"
    ws.conditional_formatting.add(f"N{first}:N{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="0F766E"))
    ws.conditional_formatting.add(f"P{first}:P{last}", FormulaRule(formula=[f'P{first}="At risk"'], fill=PatternFill("solid", fgColor=RED_FILL)))
    ws.conditional_formatting.add(f"P{first}:P{last}", FormulaRule(formula=[f'P{first}="Improvement needed"'], fill=PatternFill("solid", fgColor=AMBER_FILL)))
    ws.conditional_formatting.add(f"P{first}:P{last}", FormulaRule(formula=[f'P{first}="Preferred"'], fill=PatternFill("solid", fgColor=GREEN_FILL)))
    ws.conditional_formatting.add(f"R{first}:R{last}", FormulaRule(formula=[f'R{first}="Critical - act now"'], fill=PatternFill("solid", fgColor=RED_FILL), font=Font(bold=True)))
    note = last + 2
    ws.cell(row=note, column=1, value="How to read").font = BOLD
    notes = [
        "Each component is scaled linearly from 0 points (poor value) to 100 points (target value) and clipped; thresholds are on Assumptions.",
        "Missing components (e.g. fewer completed lines than needed for lead-time variability) are dropped and the remaining weights re-normalised - never scored as poor.",
        "Suppliers below the ranking threshold are shown as Provisional; below the provisional threshold they are not scored. Equal scores share a rank.",
        "Dependency risk is assessed separately from performance and combined only in the action quadrant.",
    ]
    for k, t in enumerate(notes, start=1):
        ws.cell(row=note + k, column=1, value=f"{k}. {t}").font = Font(color=SLATE)
    _widths(ws, {"A": 11, "B": 24, "C": 9, "D": 20, "E": 9, "F": 9, "G": 10, "H": 9, "I": 9, "J": 9, "K": 11, "L": 9,
                 "M": 9, "N": 12, "O": 9, "P": 18, "Q": 11, "R": 30})
    ws.row_dimensions[4].height = 45
    ws.freeze_panes = "C5"


def _trend(b: Book, r) -> None:
    ws = b.wb.create_sheet("Delivery_Trend")
    _title(ws, "Monthly OTIF (all synthetic suppliers)", SYN_LABEL + ". Lines grouped by promised month.")
    for j, (h, fml) in enumerate([("Promised month", False), ("Due lines", True), ("OTIF lines", True), ("OTIF rate", True),
                                  ("On-time lines", True), ("On-time rate", True)], start=1):
        _header(ws, 4, j, h, formula=fml)
    months = r.monthly_otif["month"].tolist()
    for i, m in enumerate(months, start=5):
        ws.cell(row=i, column=1, value=m.to_pydatetime()).number_format = "mmm yyyy"
        ws.cell(row=i, column=2, value=f"=COUNTIFS({b.L('due_month')},A{i},{b.L('due')},1)")
        ws.cell(row=i, column=3, value=f"=COUNTIFS({b.L('due_month')},A{i},{b.L('otif')},1)")
        ws.cell(row=i, column=4, value=f"=IF(B{i}>0,C{i}/B{i},\"\")").number_format = PCT
        ws.cell(row=i, column=5, value=f"=COUNTIFS({b.L('due_month')},A{i},{b.L('on_time')},1)")
        ws.cell(row=i, column=6, value=f"=IF(B{i}>0,E{i}/B{i},\"\")").number_format = PCT
    b.cells["trend_rows"] = (5, 4 + len(months))
    _widths(ws, {"A": 16, "B": 11, "C": 11, "D": 11, "E": 12, "F": 12})


def _exceptions(b: Book, r) -> None:
    ws = b.wb.create_sheet("Exceptions")
    _title(ws, "Supplier exception flags (live) and overdue lines", SYN_LABEL + ". Flags recalculate from Supplier_Summary and Assumptions.")
    a = b.cells
    sc = a["sum_cols"]
    heads = ["Supplier ID", "S1 Chronic late", "S2 Deteriorating", "S3 Unstable lead time", "S4 Incomplete deliveries",
             "S5 Quality", "S6 PPV above contract", "S8 High dependency", "Flags raised"]
    for j, h in enumerate(heads, start=1):
        _header(ws, 4, j, h, formula=j > 1)
    sups = r.syn.suppliers["supplier_id"].tolist()
    first = 5
    for i, sup in enumerate(sups):
        row = first + i
        srow = a["sum_first"] + i
        S = lambda k: f"Supplier_Summary!{sc[k]}{srow}"  # noqa: E731
        ws.cell(row=row, column=1, value=sup).font = BOLD
        enough = f"{S('due_lines')}>={a['x_min']}"
        ws.cell(row=row, column=2, value=f"=IF(AND({enough},{S('otif_rate')}<{a['x_late']}),\"FLAG\",\"\")")
        ws.cell(row=row, column=3, value=(f"=IF(AND({S('due_prior')}>={a['x_min']},{S('due_last')}>={a['x_min']}),"
                                          f"IF({S('otif_change')}<=-{a['x_det']},\"FLAG\",\"\"),\"\")"))
        ws.cell(row=row, column=4, value=f"=IF(AND({S('completed')}>={a['min_completed']},{S('lt_sd')}>={a['x_lt']}),\"FLAG\",\"\")")
        ws.cell(row=row, column=5, value=f"=IF(AND({enough},{S('in_full_rate')}<{a['x_if']}),\"FLAG\",\"\")")
        ws.cell(row=row, column=6, value=f"=IF(AND({enough},{S('reject_rate')}>={a['x_rej']}),\"FLAG\",\"\")")
        ws.cell(row=row, column=7, value=f"=IF({S('ppv_pct')}>={a['x_ppv']},\"FLAG\",\"\")")
        ws.cell(row=row, column=8, value=f"=IF(Scorecard!Q{a['sc_first'] + i}=\"High\",\"FLAG\",\"\")")
        ws.cell(row=row, column=9, value=f"=COUNTIF(B{row}:H{row},\"FLAG\")")
    last = first + len(sups) - 1
    b.cells["exc_rows"] = (first, last)
    ws.conditional_formatting.add(f"B{first}:H{last}", CellIsRule(operator="equal", formula=['"FLAG"'],
                                                                  fill=PatternFill("solid", fgColor=RED_FILL), font=Font(bold=True, color="991B1B")))
    row = last + 3
    ws.cell(row=row, column=1, value="S7 Overdue open PO lines (list exported by the pipeline; count check is live)").font = BOLD
    od = r.facts[r.facts["overdue"]].sort_values(["days_late", "po_line_id"], ascending=[False, True])
    ws.cell(row=row + 1, column=1, value="Live count of overdue lines:")
    ws.cell(row=row + 1, column=3, value=f"=SUM({b.L('overdue')})")
    ws.cell(row=row + 1, column=4, value=len(od))
    ws.cell(row=row + 1, column=5, value=f"=IF(C{row + 1}=D{row + 1},\"List complete\",\"Re-run pipeline\")")
    b.cells["overdue_count_cell"] = f"Exceptions!$C${row + 1}"
    hr = row + 3
    for j, h in enumerate(["PO line", "Supplier", "Item", "Promised", "Ordered", "Received", "Open value", "Days past promise"], start=1):
        _header(ws, hr, j, h)
    for k, rec in enumerate(od.itertuples(), start=hr + 1):
        vals = [rec.po_line_id, rec.supplier_id, rec.item_id, rec.promised_date.to_pydatetime(), rec.ordered_qty,
                rec.received_qty, float(rec.open_value), float(rec.days_late)]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=k, column=j, value=v)
            c.number_format = {4: DATE, 7: USD2, 8: NUM}.get(j, "General")
    _widths(ws, {"A": 16, "B": 14, "C": 14, "D": 14, "E": 14, "F": 12, "G": 14, "H": 14, "I": 10})
    ws.row_dimensions[4].height = 32


def _scenario(b: Book, r) -> None:
    ws = b.wb.create_sheet("Scenario")
    a = b.cells
    _title(ws, "What-if scenarios (modelled, not realised savings)", SYN_LABEL + " for A and B; C uses observed City of Austin prices.")
    hl = r.headline
    from_sup = hl.get("critical_supplier") or "SUP-01"
    to_sup = hl.get("realloc_to") or "SUP-02"
    cat = r.scorecard.set_index("supplier_id").loc[from_sup, "main_category_id"]
    ws["A4"] = "A. Reallocate volume between suppliers (last 12 months of due lines)"
    ws["A4"].font = BOLD
    inputs = [("Category", cat, "General"), ("Move volume from supplier", from_sup, "General"),
              ("Move volume to supplier", to_sup, "General"), ("Share of volume moved", 0.30, PCT)]
    for i, (lab, val, fmt) in enumerate(inputs, start=5):
        ws.cell(row=i, column=1, value=lab)
        c = ws.cell(row=i, column=2, value=val)
        c.fill = PatternFill("solid", fgColor=INPUT_FILL)
        c.border = BOX
        c.number_format = fmt
    ws.cell(row=9, column=1, value="Cost per late line (Assumptions)")
    ws.cell(row=9, column=2, value=f"={a['cost_late']}").number_format = USD
    dv = DataValidation(type="list", formula1=f"=Syn_Suppliers!$A${a['sup_first']}:$A${a['sup_last']}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B6")
    dv.add("B7")
    dv_cat = DataValidation(type="list", formula1='"' + ",".join(r.syn.categories["category_id"]) + '"')
    ws.add_data_validation(dv_cat)
    dv_cat.add("B5")

    L = b.L
    def vol(sup_cell: str, what: str) -> str:
        base = f"{L('supplier_id')},{sup_cell},{L('category_id')},$B$5,{L('due_period')},\"last_12m\""
        if what == "due":
            return f"COUNTIFS({base},{L('due')},1)"
        if what == "otif":
            return f"COUNTIFS({base},{L('due')},1,{L('otif')},1)"
        if what == "spend":
            return f"SUMIFS({L('line_value')},{base},{L('cancelled')},0)"
        if what == "std":
            return f"SUMIFS({L('standard_value')},{base},{L('cancelled')},0)"
        raise ValueError(what)

    _header(ws, 11, 1, "Last 12 months")
    _header(ws, 11, 2, "From supplier", formula=True)
    _header(ws, 11, 3, "To supplier", formula=True)
    rows = [("Due lines", "due", NUM), ("OTIF lines", "otif", NUM), ("Spend", "spend", USD), ("Value at contract price", "std", USD)]
    for i, (lab, what, fmt) in enumerate(rows, start=12):
        ws.cell(row=i, column=1, value=lab)
        ws.cell(row=i, column=2, value=f"={vol('$B$6', what)}").number_format = fmt
        ws.cell(row=i, column=3, value=f"={vol('$B$7', what)}").number_format = fmt
    ws.cell(row=16, column=1, value="OTIF rate")
    ws["B16"] = "=IF(B12>0,B13/B12,\"\")"
    ws["C16"] = "=IF(C12>0,C13/C12,\"\")"
    ws.cell(row=17, column=1, value="PPV vs contract")
    ws["B17"] = "=IF(B15>0,B14/B15-1,\"\")"
    ws["C17"] = "=IF(C15>0,C14/C15-1,\"\")"
    for c in ["B16", "C16", "B17", "C17"]:
        ws[c].number_format = PCT

    _header(ws, 19, 1, "Result at the chosen share")
    _header(ws, 19, 2, "Value", formula=True)
    results = [
        ("Lines moved per year", "=B8*B12", NUM1),
        ("Spend moved per year", "=B8*B14", USD),
        ("Late lines per year - before", "=B12*(1-B16)+C12*(1-C16)", NUM1),
        ("Late lines per year - after", "=(1-B8)*B12*(1-B16)+(C12+B8*B12)*(1-C16)", NUM1),
        ("Late lines avoided per year", "=B22-B23", NUM1),
        ("Modelled late-handling cost avoided", "=B24*B9", USD),
        ("Modelled price impact (+ costs more)", "=B21*((1+C17)/(1+B17)-1)", USD),
        ("Net modelled benefit per year", "=B25-B26", USD),
    ]
    for i, (lab, fml, fmt) in enumerate(results, start=20):
        ws.cell(row=i, column=1, value=lab)
        ws.cell(row=i, column=2, value=fml).number_format = fmt
    b.cells["scen_avoided"] = "Scenario!$B$24"
    b.cells["scen_price"] = "Scenario!$B$26"
    ws["A28"] = "Assumes the receiving supplier keeps its current OTIF and pricing at higher volume; confirm capacity and contract terms first."
    ws["A28"].font = Font(color=SLATE, italic=True)

    _header(ws, 30, 1, "Sensitivity: share moved")
    for j, h in enumerate(["Late lines avoided", "Cost avoided", "Price impact", "Net benefit"], start=2):
        _header(ws, 30, j, h, formula=True)
    for i, share in enumerate(r.settings.scenario["reallocation_shares"], start=31):
        ws.cell(row=i, column=1, value=share).number_format = PCT
        ws.cell(row=i, column=2, value=f"=A{i}*B12*(C16-B16)").number_format = NUM1
        ws.cell(row=i, column=3, value=f"=B{i}*$B$9").number_format = USD
        ws.cell(row=i, column=4, value=f"=A{i}*$B$14*((1+$C$17)/(1+$B$17)-1)").number_format = USD
        ws.cell(row=i, column=5, value=f"=C{i}-D{i}").number_format = USD

    start = 39
    ws.cell(row=start, column=1, value="B. Late lines avoided if each ranked supplier reached the OTIF target").font = BOLD
    for j, (h, fml) in enumerate([("Supplier", False), ("Due lines last 12m", True), ("OTIF last 12m", True), ("Late lines now", True),
                                  ("Late lines at target", True), ("Late lines avoided", True), ("Cost avoided", True)], start=1):
        _header(ws, start + 1, j, h, formula=fml)
    sc = a["sum_cols"]
    ranked = r.otif_scenario["supplier_id"].tolist()
    order = r.syn.suppliers["supplier_id"].tolist()
    for k, sup in enumerate(ranked, start=start + 2):
        srow = a["sum_first"] + order.index(sup)
        ws.cell(row=k, column=1, value=sup)
        ws.cell(row=k, column=2, value=f"=Supplier_Summary!{sc['due_last']}{srow}").number_format = NUM
        ws.cell(row=k, column=3, value=f"=IF(Supplier_Summary!{sc['otif_last']}{srow}=\"\",Supplier_Summary!{sc['otif_rate']}{srow},Supplier_Summary!{sc['otif_last']}{srow})").number_format = PCT
        ws.cell(row=k, column=4, value=f"=B{k}*(1-C{k})").number_format = NUM1
        ws.cell(row=k, column=5, value=f"=B{k}*(1-MAX(C{k},{a['otif_t']}))").number_format = NUM1
        ws.cell(row=k, column=6, value=f"=D{k}-E{k}").number_format = NUM1
        ws.cell(row=k, column=7, value=f"=F{k}*{a['cost_late']}").number_format = USD
    b.cells["otif_scen_rows"] = (start + 2, start + 1 + len(ranked))

    pstart = start + 3 + len(ranked)
    ws.cell(row=pstart, column=1, value="C. Price harmonisation on like-for-like items (OBSERVED City of Austin prices)").font = BOLD
    ws.cell(row=pstart + 1, column=1, value=f"Excess paid above item-year median on flagged lines, {r.headline['period_label']} total (pipeline value)")
    c = ws.cell(row=pstart + 1, column=2, value=round(r.prices.summary["modelled_opportunity_full"], 2))
    c.number_format = USD
    _header(ws, pstart + 3, 1, "Capture rate (input)")
    _header(ws, pstart + 3, 2, "Modelled opportunity", formula=True)
    for i, rate in enumerate(r.settings.scenario["price_capture_rates"], start=pstart + 4):
        c = ws.cell(row=i, column=1, value=rate)
        c.number_format = PCT
        c.fill = PatternFill("solid", fgColor=INPUT_FILL)
        ws.cell(row=i, column=2, value=f"=A{i}*$B${pstart + 1}").number_format = USD
    ws.cell(row=pstart + 8, column=1, value="Not a savings forecast: price differences may reflect urgency, order size, delivery terms or market timing.").font = Font(color=SLATE, italic=True)
    _widths(ws, {"A": 44, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 14})


def _austin_vendors(b: Book, r, period: str) -> None:
    from .pipeline import public_vendor_table

    ws = b.wb.create_sheet("Austin_Vendors")
    _title(ws, "Vendor spend and concentration", OBS_LABEL.format(period=period) + ". Top 150 vendors; the rest aggregated.")
    v = public_vendor_table(r.vendors)
    heads = [("Vendor code", False), ("Vendor name", False), ("Spend (USD)", False), ("POs", False), ("NIGP classes", False),
             ("Spend on master agreements", False), ("Share of total", True), ("Cumulative share", True), ("Rank", True),
             ("Pareto class", True), ("Avg PO value", True), ("Agreement share", True)]
    for j, (h, fml) in enumerate(heads, start=1):
        _header(ws, 4, j, h, formula=fml)
    first = 5
    last = first + len(v) - 1
    top_last = last - 1 if v.iloc[-1]["vendor_code"] == "ALL_OTHER" else last
    for i, rec in enumerate(v.itertuples(), start=first):
        ws.cell(row=i, column=1, value=rec.vendor_code)
        ws.cell(row=i, column=2, value=rec.vendor_name)
        ws.cell(row=i, column=3, value=round(float(rec.spend), 2)).number_format = USD
        ws.cell(row=i, column=4, value=int(rec.po_count)).number_format = NUM
        ws.cell(row=i, column=5, value=None if pd.isna(rec.categories) else int(rec.categories))
        ws.cell(row=i, column=6, value=round(float(rec.agreement_spend), 2)).number_format = USD
        ws.cell(row=i, column=7, value=f"=C{i}/SUM($C${first}:$C${last})").number_format = PCT
        if i <= top_last:
            ws.cell(row=i, column=8, value=f"=SUM($C${first}:C{i})/SUM($C${first}:$C${last})").number_format = PCT
            ws.cell(row=i, column=9, value=f"=RANK.EQ(C{i},$C${first}:$C${top_last},0)")
            ws.cell(row=i, column=10, value=f"=IF(H{i}-G{i}<0.8,\"A\",IF(H{i}-G{i}<0.95,\"B\",\"C\"))")
        ws.cell(row=i, column=11, value=f"=C{i}/D{i}").number_format = USD
        ws.cell(row=i, column=12, value=f"=IF(C{i}>0,F{i}/C{i},\"\")").number_format = PCT
    tot = last + 1
    ws.cell(row=tot, column=2, value="Total").font = BOLD
    ws.cell(row=tot, column=3, value=f"=SUM(C{first}:C{last})").number_format = USD
    ws.cell(row=tot, column=4, value=f"=SUM(D{first}:D{last})").number_format = NUM
    ws.cell(row=tot, column=6, value=f"=SUM(F{first}:F{last})").number_format = USD
    ws.cell(row=tot, column=12, value=f"=F{tot}/C{tot}").number_format = PCT
    b.cells["austin_total"] = f"Austin_Vendors!$C${tot}"
    b.cells["austin_agreement_share"] = f"Austin_Vendors!$L${tot}"
    ws.cell(row=tot + 2, column=2, value="Top-10 vendor share (formula)").font = BOLD
    ws.cell(row=tot + 2, column=3, value=f"=SUMIFS(G{first}:G{top_last},I{first}:I{top_last},\"<=10\")").number_format = PCT
    b.cells["austin_top10"] = f"Austin_Vendors!$C${tot + 2}"
    ws.cell(row=tot + 3, column=2, value="Vendors to reach 80% of spend (formula)").font = BOLD
    ws.cell(row=tot + 3, column=3, value=f"=COUNTIFS(J{first}:J{top_last},\"A\")")
    b.cells["austin_v80"] = f"Austin_Vendors!$C${tot + 3}"
    ws.cell(row=tot + 4, column=2, value="Vendor HHI, all 1,900+ vendors (pipeline value)").font = BOLD
    ws.cell(row=tot + 4, column=3, value=round(r.kpis["vendor_hhi"], 1))
    ws.conditional_formatting.add(f"G{first}:G{top_last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="1F2A44"))
    tab = Table(displayName="tblAustinVendors", ref=f"A4:L{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "C5"
    _widths(ws, {"A": 15, "B": 42, "C": 16, "D": 9, "E": 9, "F": 18, "G": 10, "H": 11, "I": 7, "J": 8, "K": 13, "L": 11})
    ws.row_dimensions[4].height = 32


def _austin_categories(b: Book, r, period: str) -> None:
    ws = b.wb.create_sheet("Austin_Categories")
    c = r.settings.concentration
    _title(ws, "Category spend and supplier dependency", OBS_LABEL.format(period=period) + f". NIGP classes with spend >= ${c['min_category_spend']:,.0f}.")
    ws["A3"] = "Thresholds (inputs):"
    ws["A3"].font = BOLD
    ws["C3"] = "Single-vendor share"
    ws["D3"] = c["single_vendor_share"]
    ws["E3"] = "HHI highly concentrated"
    ws["F3"] = c["hhi_highly_concentrated"]
    for cell, fmt in [("D3", PCT), ("F3", NUM)]:
        ws[cell].fill = PatternFill("solid", fgColor=INPUT_FILL)
        ws[cell].number_format = fmt
    cats = r.categories[r.categories["spend"] >= c["min_category_spend"]].reset_index(drop=True)
    heads = [("NIGP class", False), ("Class label", False), ("Spend family", False), ("Spend (USD)", False), ("Vendors", False),
             ("Top vendor", False), ("Top vendor share", False), ("Top vendor orders", False), ("HHI", False),
             ("Standalone spend (no master agreement)", False), ("Share of total", True), ("Spend rank", True), ("Dependency flag", True),
             ("Agreement coverage", True)]
    for j, (h, fml) in enumerate(heads, start=1):
        _header(ws, 5, j, h, formula=fml)
    first = 6
    last = first + len(cats) - 1
    for i, rec in enumerate(cats.itertuples(), start=first):
        vals = [rec.nigp_class, rec.class_label, rec.spend_family, round(rec.spend, 2), int(rec.vendor_count), rec.top_vendor_name,
                float(rec.top_vendor_share), int(rec.top_vendor_orders), round(float(rec.hhi), 1), round(float(rec.spend_without_agreement), 2)]
        fmts = [None, None, None, USD, NUM, None, PCT, NUM, NUM, USD]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            cell = ws.cell(row=i, column=j, value=v)
            if f:
                cell.number_format = f
        ws.cell(row=i, column=11, value=f"=D{i}/{b.cells['austin_total']}").number_format = PCT
        ws.cell(row=i, column=12, value=f"=RANK.EQ(D{i},$D${first}:$D${last},0)")
        ws.cell(row=i, column=13, value=f"=IF(G{i}>=$D$3,\"R1 single-vendor\",IF(I{i}>=$F$3,\"R2 highly concentrated\",\"\"))")
        ws.cell(row=i, column=14, value=f"=1-J{i}/D{i}").number_format = PCT
    ws.conditional_formatting.add(f"M{first}:M{last}", FormulaRule(formula=[f'LEFT(M{first},2)="R1"'], fill=PatternFill("solid", fgColor=RED_FILL), font=Font(bold=True)))
    ws.conditional_formatting.add(f"M{first}:M{last}", FormulaRule(formula=[f'LEFT(M{first},2)="R2"'], fill=PatternFill("solid", fgColor=AMBER_FILL)))
    ws.conditional_formatting.add(f"I{first}:I{last}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=10000, end_color="9FB4CC"))
    s = last + 2
    ws.cell(row=s, column=2, value="Summary (formulas)").font = BOLD
    summ = [("Categories screened", f"=COUNTA(A{first}:A{last})", NUM),
            ("R1 single-vendor categories", f"=COUNTIF(M{first}:M{last},\"R1*\")", NUM),
            ("R2 highly concentrated categories", f"=COUNTIF(M{first}:M{last},\"R2*\")", NUM),
            ("Spend in R1 categories", f"=SUMIFS(D{first}:D{last},M{first}:M{last},\"R1*\")", USD),
            ("Largest category share held by one vendor", f"=MAXIFS(G{first}:G{last},D{first}:D{last},\">=\"&10000000)", PCT)]
    for k, (lab, fml, fmt) in enumerate(summ, start=s + 1):
        ws.cell(row=k, column=2, value=lab)
        ws.cell(row=k, column=4, value=fml).number_format = fmt
    b.cells["austin_r1_count"] = f"Austin_Categories!$D${s + 2}"
    b.cells["austin_r1_spend"] = f"Austin_Categories!$D${s + 4}"
    tab = Table(displayName="tblAustinCategories", ref=f"A5:N{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "C6"
    _widths(ws, {"A": 8, "B": 46, "C": 26, "D": 15, "E": 8, "F": 30, "G": 10, "H": 9, "I": 8, "J": 15, "K": 9, "L": 7, "M": 20, "N": 10})
    ws.row_dimensions[5].height = 45

    # Family summary
    fs = b.wb.create_sheet("Austin_Families")
    _title(fs, "Spend by family and fiscal year", OBS_LABEL.format(period=period) + ". Families are an analyst grouping of NIGP classes.")
    ft = r.family_trend
    years = list(ft.columns)
    _header(fs, 4, 1, "Spend family")
    for j, y in enumerate(years, start=2):
        _header(fs, 4, j, f"FY{y}")
    tcol = len(years) + 2
    _header(fs, 4, tcol, "Total", formula=True)
    _header(fs, 4, tcol + 1, "Share", formula=True)
    _header(fs, 4, tcol + 2, f"Change FY{years[0]}-FY{years[-1]}", formula=True)
    ft = ft.loc[ft.sum(axis=1).sort_values(ascending=False).index]
    for i, (fam, row) in enumerate(ft.iterrows(), start=5):
        fs.cell(row=i, column=1, value=fam)
        for j, y in enumerate(years, start=2):
            fs.cell(row=i, column=j, value=round(float(row[y]), 2)).number_format = USD
        lc = get_column_letter(tcol - 1)
        fs.cell(row=i, column=tcol, value=f"=SUM(B{i}:{lc}{i})").number_format = USD
        tl = get_column_letter(tcol)
        fs.cell(row=i, column=tcol + 1, value=f"={tl}{i}/SUM(${tl}$5:${tl}${4 + len(ft)})").number_format = PCT
        fs.cell(row=i, column=tcol + 2, value=f"=IF(B{i}>0,{lc}{i}/B{i}-1,\"\")").number_format = PCT
    tot = 5 + len(ft)
    fs.cell(row=tot, column=1, value="Total").font = BOLD
    for j in range(2, tcol + 1):
        col = get_column_letter(j)
        fs.cell(row=tot, column=j, value=f"=SUM({col}5:{col}{tot - 1})").number_format = USD
    b.cells["family_total"] = f"Austin_Families!${get_column_letter(tcol)}${tot}"
    fs.column_dimensions["A"].width = 34
    for j in range(2, tcol + 3):
        fs.column_dimensions[get_column_letter(j)].width = 15


def _austin_prices(b: Book, r, period: str) -> None:
    ws = b.wb.create_sheet("Austin_Prices")
    p = r.settings.price
    _title(ws, "Like-for-like purchase price variance", OBS_LABEL.format(period=period)
           + ". Same commodity code + unit + description + fiscal year; >= 4 purchases per item-year.")
    ws["A3"] = "Flag ratio (input):"
    ws["C3"] = p["above_baseline_ratio"]
    ws["C3"].fill = PatternFill("solid", fgColor=INPUT_FILL)
    iy = r.prices.item_years[r.prices.item_years["flagged_lines"] > 0].head(60).reset_index(drop=True)
    heads = [("Item description (redacted)", False), ("Unit", False), ("FY", False), ("Lines", False), ("Vendors", False),
             ("Item-year median price", False), ("Max price", False), ("Flagged lines", False), ("Excess vs median (USD)", False),
             ("Max / median", True), ("Above flag ratio?", True)]
    for j, (h, fml) in enumerate(heads, start=1):
        _header(ws, 5, j, h, formula=fml)
    first = 6
    for i, rec in enumerate(iy.itertuples(), start=first):
        vals = [rec.item_description[:120], rec.uom, int(rec.fiscal_year), int(rec.lines), int(rec.vendors),
                float(rec.baseline_price), float(rec.max_price), int(rec.flagged_lines), round(float(rec.excess_vs_baseline), 2)]
        fmts = [None, None, "0", NUM, NUM, USD2, USD2, NUM, USD]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            c = ws.cell(row=i, column=j, value=v)
            if f:
                c.number_format = f
        ws.cell(row=i, column=10, value=f"=G{i}/F{i}").number_format = "0.00x"
        ws.cell(row=i, column=11, value=f"=IF(J{i}>=$C$3,\"Yes\",\"No\")")
    last = first + len(iy) - 1
    ws.cell(row=last + 1, column=1, value=f"Top {len(iy)} of {int((r.prices.item_years['flagged_lines'] > 0).sum())} flagged item-years shown; full list in outputs/tables/obs_price_item_years.csv").font = Font(color=SLATE, italic=True)

    y = r.prices.yoy[r.prices.yoy["pct_change"] >= 0.25].reset_index(drop=True)
    s = last + 4
    ws.cell(row=s - 1, column=1, value="Year-over-year median price increases of 25% or more (same item, consecutive fiscal years)").font = BOLD
    yh = [("Item description (redacted)", False), ("Unit", False), ("From FY", False), ("To FY", False), ("Prior median", False),
          ("New median", False), ("Spend in later year", False), ("Change", True)]
    for j, (h, fml) in enumerate(yh, start=1):
        _header(ws, s, j, h, formula=fml)
    for i, rec in enumerate(y.itertuples(), start=s + 1):
        vals = [rec.item_description[:120], rec.uom, int(rec.prior_fiscal_year), int(rec.fiscal_year), float(rec.prior_baseline),
                float(rec.baseline_price), round(float(rec.spend), 2)]
        fmts = [None, None, "0", "0", USD2, USD2, USD]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            c = ws.cell(row=i, column=j, value=v)
            if f:
                c.number_format = f
        ws.cell(row=i, column=8, value=f"=F{i}/E{i}-1").number_format = PCT
    ylast = s + len(y)
    ws.conditional_formatting.add(f"H{s + 1}:H{ylast}", DataBarRule(start_type="num", start_value=0, end_type="max", color="B45309"))
    ws.cell(row=ylast + 2, column=1, value="Median change across these items (formula):").font = BOLD
    ws.cell(row=ylast + 2, column=8, value=f"=MEDIAN(H{s + 1}:H{ylast})").number_format = PCT
    _widths(ws, {"A": 70, "B": 7, "C": 8, "D": 8, "E": 13, "F": 14, "G": 14, "H": 10, "I": 16, "J": 10, "K": 10})
    ws.row_dimensions[5].height = 32


def _actions(b: Book, r) -> None:
    ws = b.wb.create_sheet("Action_List")
    _title(ws, "Prioritised procurement actions", "Each row cites its data origin. Synthetic rows illustrate the method; observed rows come from City of Austin data.")
    cols = ["action_id", "priority", "data_origin", "supplier_or_category", "issue", "supporting_metric", "business_significance",
            "recommended_action", "data_limitation", "source_rules"]
    heads = ["ID", "Priority", "Data origin", "Supplier or category", "Issue", "Supporting metric", "Business significance",
             "Recommended action", "Data limitation", "Exception IDs"]
    for j, h in enumerate(heads, start=1):
        _header(ws, 4, j, h)
    for i, rec in enumerate(r.actions[cols].itertuples(index=False), start=5):
        for j, v in enumerate(rec, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    last = 4 + len(r.actions)
    ws.conditional_formatting.add(f"B5:B{last}", CellIsRule(operator="equal", formula=['"High"'], fill=PatternFill("solid", fgColor=RED_FILL), font=Font(bold=True)))
    ws.conditional_formatting.add(f"B5:B{last}", CellIsRule(operator="equal", formula=['"Medium"'], fill=PatternFill("solid", fgColor=AMBER_FILL)))
    ws.cell(row=last + 2, column=1, value="Count by priority (formulas):").font = BOLD
    for k, pr in enumerate(["High", "Medium", "Low"]):
        ws.cell(row=last + 3 + k, column=1, value=pr)
        ws.cell(row=last + 3 + k, column=2, value=f"=COUNTIF(B5:B{last},\"{pr}\")")
    _widths(ws, {"A": 6, "B": 9, "C": 16, "D": 26, "E": 24, "F": 38, "G": 34, "H": 44, "I": 36, "J": 16})
    ws.freeze_panes = "C5"


def _dictionary(b: Book, r) -> None:
    from .docs_text import DATA_DICTIONARY

    ws = b.wb.create_sheet("Data_Dictionary")
    _title(ws, "Data dictionary", "Fields used in this workbook and the SQL database. Origin: O = observed City of Austin, S = synthetic, D = derived.")
    for j, h in enumerate(["Table / sheet", "Field", "Origin", "Definition"], start=1):
        _header(ws, 4, j, h)
    for i, (t, f, o, d) in enumerate(DATA_DICTIONARY, start=5):
        for j, v in enumerate([t, f, o, d], start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, {"A": 22, "B": 26, "C": 8, "D": 100})


def _validation(b: Book, r) -> None:
    ws = b.wb.create_sheet("Validation")
    _title(ws, "Excel vs Python reconciliation", "Python values are fixed at export; Excel values are live formulas. The test suite recalculates this sheet independently.")
    a = b.cells
    for j, h in enumerate(["Check", "Excel (formula)", "Python pipeline", "Difference", "Result"], start=1):
        _header(ws, 4, j, h, formula=j in (2, 4, 5))
    sc = r.scorecard.set_index("supplier_id")
    order = r.syn.suppliers["supplier_id"].tolist()
    checks: list[tuple[str, str, float]] = []
    tot = a["sum_total_row"]
    cols = a["sum_cols"]
    checks.append(("Synthetic: overall OTIF rate", f"=Supplier_Summary!{cols['otif_rate']}{tot}", r.syn_kpis["otif_rate"]))
    checks.append(("Synthetic: total spend", f"=Supplier_Summary!{cols['spend']}{tot}", r.syn_kpis["spend"]))
    checks.append(("Synthetic: overdue lines", f"={a['overdue_count_cell']}", r.syn_kpis["overdue_lines"]))
    checks.append(("Synthetic: reject rate", f"=Supplier_Summary!{cols['reject_rate']}{tot}", r.syn_kpis["reject_rate"]))
    for sup in order:
        i = order.index(sup)
        score = sc.loc[sup, "performance_score"]
        checks.append((f"Score {sup}", f"=IF(Scorecard!N{a['sc_first'] + i}=\"\",-1,Scorecard!N{a['sc_first'] + i})",
                       -1.0 if pd.isna(score) else float(score)))
    for sup in order:
        i = order.index(sup)
        rank = sc.loc[sup, "rank"]
        checks.append((f"Rank {sup}", f"=IF(Scorecard!O{a['sc_first'] + i}=\"\",-1,Scorecard!O{a['sc_first'] + i})",
                       -1.0 if pd.isna(rank) else float(rank)))
    for sup in order:
        i = order.index(sup)
        checks.append((f"OTIF {sup}", f"=Supplier_Summary!{cols['otif_rate']}{a['sum_first'] + i}", float(sc.loc[sup, "otif_rate"])))
    crit = r.headline.get("critical_supplier")
    rr = r.realloc[r.realloc["share_moved"].round(2) == 0.30] if len(r.realloc) else pd.DataFrame()
    if len(rr):
        checks.append((f"Scenario: late lines avoided (30% {crit} -> {r.headline['realloc_to']})", f"={a['scen_avoided']}", float(rr.iloc[0]["late_lines_avoided"])))
        checks.append(("Scenario: price impact at 30%", f"={a['scen_price']}", float(rr.iloc[0]["price_impact"])))
    checks.append(("Observed: total spend", f"={a['austin_total']}", r.kpis["total_spend"]))
    checks.append(("Observed: master-agreement share", f"={a['austin_agreement_share']}", r.kpis["master_agreement_share"]))
    checks.append(("Observed: top-10 vendor share", f"={a['austin_top10']}", r.kpis["top10_vendor_share"]))
    checks.append(("Observed: vendors for 80% of spend", f"={a['austin_v80']}", r.kpis["vendors_for_80pct"]))
    checks.append(("Observed: family totals = total spend", f"={a['family_total']}", r.kpis["total_spend"]))
    checks.append(("Observed: single-vendor (R1) categories", f"={a['austin_r1_count']}", r.headline["single_vendor_categories"]))
    checks.append(("Observed: spend in R1 categories", f"={a['austin_r1_spend']}", r.headline["single_vendor_spend"]))
    for i, (label, fml, py) in enumerate(checks, start=5):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=fml).number_format = "#,##0.0000"
        ws.cell(row=i, column=3, value=float(py)).number_format = "#,##0.0000"
        ws.cell(row=i, column=4, value=f"=ABS(B{i}-C{i})").number_format = "0.000000"
        ws.cell(row=i, column=5, value=f"=IF(D{i}<=MAX(0.0005,ABS(C{i})*0.000001),\"MATCH\",\"CHECK\")")
    last = 4 + len(checks)
    b.cells["validation_rows"] = (5, last)
    ws.cell(row=last + 2, column=1, value="Checks matching").font = BOLD
    ws.cell(row=last + 2, column=2, value=f"=COUNTIF(E5:E{last},\"MATCH\")&\" of \"&COUNTA(A5:A{last})")
    ws.conditional_formatting.add(f"E5:E{last}", CellIsRule(operator="equal", formula=['"MATCH"'], fill=PatternFill("solid", fgColor=GREEN_FILL)))
    ws.conditional_formatting.add(f"E5:E{last}", CellIsRule(operator="equal", formula=['"CHECK"'], fill=PatternFill("solid", fgColor=RED_FILL)))
    _widths(ws, {"A": 52, "B": 18, "C": 18, "D": 14, "E": 10})


def _instructions(ws, b: Book, r, period: str) -> None:
    hl = r.headline
    _title(ws, "SupplierPulse - Supplier Performance & Procurement Analytics", "Procurement analysis workbook - author: Harjot Sandhu")
    lines = [
        ("What this workbook does", None),
        ("Evaluates supplier reliability, purchasing spend and procurement risk, and turns them into a prioritised action list.", None),
        ("Two data sources - never mixed", None),
        (f"OBSERVED: City of Austin purchase-order lines, {period} ({hl['period_dates']}), from data.austintexas.gov. "
         "Supports spend, concentration, agreement coverage and like-for-like price analysis. It has no delivery dates.", None),
        ("SYNTHETIC: a deterministic demonstration dataset of 14 anonymous suppliers (SUP-01 ... SUP-14), purchase orders and goods "
         "receipts. Used only to demonstrate delivery, quality, scorecard and scenario methods. Not real suppliers or real performance.", None),
        ("Header colours", None),
        ("Navy header = data loaded from the Python pipeline (values).   Teal header = live Excel formula.", None),
        ("Sheets", None),
    ]
    sheets = [
        ("Assumptions", "Inputs and thresholds used by every formula (yellow cells)."),
        ("Syn_Lines", "Synthetic PO lines (Excel table tblLines): loaded line facts + formula columns for due, on-time, in-full, OTIF, days late, lead time."),
        ("Syn_Suppliers", "Synthetic supplier master and the designed behaviour of each supplier."),
        ("Supplier_Summary", "Per-supplier delivery, quality, spend and PPV metrics (COUNTIFS / SUMIFS / AVERAGEIFS)."),
        ("Dependency", "Supplier spend share by category, main category and dependency tier."),
        ("Scorecard", "Weighted 0-100 performance score, evidence status, competition rank, tier and action quadrant (XLOOKUP for names)."),
        ("Delivery_Trend", "Monthly OTIF by promised month."),
        ("Exceptions", "Live supplier exception flags S1-S8 and the overdue-line list."),
        ("Scenario", "A: volume reallocation model with inputs and sensitivity; B: OTIF-target scenario; C: observed price harmonisation."),
        ("Austin_Vendors", "Observed vendor spend, share, cumulative share, rank and Pareto class (top 150 + aggregate)."),
        ("Austin_Categories", "Observed category spend, top-vendor share, HHI and dependency flags."),
        ("Austin_Families", "Observed spend by family and fiscal year."),
        ("Austin_Prices", "Observed like-for-like price variance and year-over-year price changes."),
        ("Action_List", "Prioritised actions with metric, significance, action and data limitation."),
        ("Data_Dictionary", "Field definitions."),
        ("Validation", "Excel formulas reconciled against the Python pipeline."),
    ]
    row = 4
    for text, _ in lines:
        c = ws.cell(row=row, column=1, value=text)
        if text in ("What this workbook does", "Two data sources - never mixed", "Header colours", "Sheets"):
            c.font = BOLD
            row += 1
        else:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            ws.row_dimensions[row].height = 32
            row += 1
    for name, desc in sheets:
        ws.cell(row=row, column=1, value=name).font = Font(bold=True, color=TEAL)
        ws.cell(row=row, column=2, value=desc).alignment = Alignment(wrap_text=True)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Testing note").font = BOLD
    ws.cell(row=row + 1, column=1, value=("Formulas calculate when the file is opened (full recalculation on load). This workbook was generated with "
                                          "openpyxl and verified by recalculating it with the Python 'formulas' engine and comparing results with "
                                          "pandas and SQL; it has not been manually opened in Microsoft Excel by the author's build process."))
    ws.cell(row=row + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=3)
    ws.row_dimensions[row + 1].height = 48
    ws.cell(row=row + 3, column=1, value="Modelled values (scenarios, price opportunity) are estimates under stated assumptions, not realised savings.").font = Font(italic=True, color=SLATE)
    _widths(ws, {"A": 22, "B": 100, "C": 4})
