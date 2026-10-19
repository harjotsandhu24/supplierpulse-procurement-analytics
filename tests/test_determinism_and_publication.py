"""Determinism, public/private separation and privacy safeguards."""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from supplierpulse import privacy_scan, synthetic
from supplierpulse.ingest import SNAPSHOT_COLUMNS, load_snapshot, reviewed_item_texts, reviewed_vendor_names

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "austin_po_lines_raw.csv"


def test_synthetic_generation_is_deterministic(settings):
    a, b = synthetic.generate(settings), synthetic.generate(settings)
    for name in ["suppliers", "purchase_orders", "po_lines", "receipts"]:
        pd.testing.assert_frame_equal(getattr(a, name), getattr(b, name))


def test_committed_synthetic_files_match_generator(settings):
    gen = synthetic.generate(settings)
    folder = settings.path("public_dir") / "synthetic"
    for name in ["po_lines", "receipts", "purchase_orders"]:
        disk = pd.read_csv(folder / f"syn_{name}.csv")
        assert len(disk) == len(getattr(gen, name)), name
        assert disk.iloc[:, 0].tolist() == getattr(gen, name).iloc[:, 0].tolist(), name


def test_dashboard_and_brief_are_deterministic(results, tmp_path):
    from supplierpulse import brief, dashboard

    p1 = dashboard.build_dashboard(results, tmp_path / "a.html")
    p2 = dashboard.build_dashboard(results, tmp_path / "b.html")
    assert p1.read_bytes() == p2.read_bytes()
    b1 = brief.build_brief(results, tmp_path / "a_brief.html")
    b2 = brief.build_brief(results, tmp_path / "b_brief.html")
    assert b1.read_bytes() == b2.read_bytes()


def test_gitignore_keeps_raw_and_working_data_private():
    gi = (ROOT / ".gitignore").read_text(encoding="utf-8")
    for entry in ["data/raw/", "data/processed/", ".venv/", "private/"]:
        assert entry in gi


# ---------------------------------------------------------------- privacy: structural allowlists

def test_public_snapshot_is_exactly_the_column_allowlist(settings):
    snap = load_snapshot(settings)
    assert list(snap.columns) == SNAPSHOT_COLUMNS
    for forbidden in ["LGL_NM", "AD_LN_1", "AD_LN_2", "CITY", "ZIP", "ST", "EXTENDED_DESCRIPTION", "CONTRACT_NAME",
                      "COMMODITY_DESCRIPTION", "MASTER_AGREEMENT", "VENDOR_CODE"]:
        assert forbidden not in snap.columns
    assert snap["VENDOR_ID"].str.fullmatch(r"V[0-9A-F]{10}").all()


def test_snapshot_item_keys_are_only_reviewed_hashes(settings):
    snap = load_snapshot(settings)
    keys = set(snap["ITEM_TEXT_KEY"]) - {""}
    assert keys <= set(reviewed_item_texts(settings)["text_sha256"])


def test_published_vendor_names_come_only_from_the_reviewed_list(settings, results):
    reviewed = set(reviewed_vendor_names(settings)["vendor_name"])
    withheld = set(pd.read_csv(settings.root / "config" / "withheld_vendor_ids.csv", dtype=str)["vendor_id"])
    v = results.model.vendors
    named = v[v["name_published"]]
    assert set(named["vendor_name"]) <= reviewed
    assert (v.loc[~v["name_published"], "vendor_name"] == "Vendor " + v.loc[~v["name_published"], "vendor_code"]).all()
    assert not (set(named["vendor_code"]) & withheld)
    tables = settings.path("tables_dir")
    allowed = reviewed | set("Vendor " + v["vendor_code"]) | {n for n in pd.read_csv(tables / "obs_vendor_spend_top150.csv")["vendor_name"] if n.startswith("All other vendors")}
    for f, col in [("obs_vendor_spend_top150.csv", "vendor_name"), ("obs_category_concentration.csv", "top_vendor_name"),
                   ("obs_vendor_price_index.csv", "vendor_name")]:
        assert set(pd.read_csv(tables / f)[col].dropna()) <= allowed, f


def test_published_item_descriptions_come_only_from_reviewed_labels(settings):
    labels = set(reviewed_item_texts(settings)["public_label"])
    tables = settings.path("tables_dir")
    for f in ["obs_price_item_years.csv", "obs_price_year_over_year.csv"]:
        assert set(pd.read_csv(tables / f)["item_description"].dropna()) <= labels, f
    reg = pd.read_csv(tables / "exception_register.csv")
    items = reg[reg["entity_type"].isin(["Item-year", "Item"])]["entity_name"]
    assert all(any(lab.startswith(n) for lab in labels) for n in items)


def test_workbook_text_cells_use_only_reviewed_names_and_labels(settings, results, tmp_path):
    from supplierpulse import excel_export

    path = excel_export.build_workbook(results, tmp_path / "wb.xlsx")
    wb = load_workbook(path, read_only=True)
    labels = set(reviewed_item_texts(settings)["public_label"])
    names = set(reviewed_vendor_names(settings)["vendor_name"])
    prices = [r[0] for r in wb["Austin_Prices"].iter_rows(min_row=6, values_only=True) if isinstance(r[0], str)]
    item_cells = [x for x in prices if not x.startswith(("Top ", "Item description", "Year-over-year", "Median change"))]
    assert item_cells and all(any(lab.startswith(x) for lab in labels) for x in item_cells)
    vend = [r[1] for r in wb["Austin_Vendors"].iter_rows(min_row=5, values_only=True) if isinstance(r[1], str)]
    assert all(x in names or x.startswith(("Vendor V", "All other vendors", "Total", "Top-10", "Vendors to", "Vendor HHI")) for x in vend)


# ---------------------------------------------------------------- privacy: scans

def test_no_email_or_phone_patterns_in_publishable_files():
    res = privacy_scan.pattern_scan(privacy_scan.publishable_files())
    assert res.files_scanned > 50
    assert not res.findings, res.findings[:5]


@pytest.mark.skipif(not RAW.exists(), reason="private raw file not present (clean clone)")
def test_leakage_scan_detects_a_planted_canary_and_finds_nothing_real(tmp_path):
    from supplierpulse.ingest import read_raw

    raw = read_raw(RAW)
    needles = privacy_scan.private_needles(raw, set(pd.read_csv(ROOT / "config" / "reviewed_vendor_names.csv")["vendor_name"]))
    canary_name = sorted(needles["person_name_in_free_text"])[0]
    canary_addr = sorted(needles["vendor_address"])[0]
    planted = tmp_path / "planted.csv"
    planted.write_text(f"note\n{canary_name.title()} at {canary_addr.lower()}\n", encoding="utf-8")
    original_root = privacy_scan.ROOT
    try:
        privacy_scan.ROOT = tmp_path
        hits = privacy_scan.leakage_scan([planted], needles)
    finally:
        privacy_scan.ROOT = original_root
    assert {k for _, k, _ in hits.findings} >= {"person_name_in_free_text", "vendor_address"}
    real = privacy_scan.leakage_scan(privacy_scan.publishable_files(), needles)
    assert not real.findings, [(f, k) for f, k, _ in real.findings][:5]


def test_publishable_files_are_under_github_size_limits():
    big = [(str(p), p.stat().st_size) for p in privacy_scan.publishable_files() if p.stat().st_size > 50 * 1024 * 1024]
    assert not big, big


def test_published_vendor_table_aggregates_the_long_tail(settings):
    v = pd.read_csv(settings.path("tables_dir") / "obs_vendor_spend_top150.csv")
    assert len(v) == 151 and v.iloc[-1]["vendor_code"] == "ALL_OTHER"
