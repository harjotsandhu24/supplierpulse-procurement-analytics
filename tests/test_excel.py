"""Excel workbook: genuine formulas, and (slow) independent recalculation that must match Python."""

import re

import pytest
from openpyxl import load_workbook

from supplierpulse import excel_export


@pytest.fixture(scope="module")
def workbook_path(results, tmp_path_factory):
    return excel_export.build_workbook(results, tmp_path_factory.mktemp("xl") / "SupplierPulse.xlsx")


def test_expected_sheets(workbook_path):
    wb = load_workbook(workbook_path)
    for name in ["Instructions", "Assumptions", "Syn_Lines", "Supplier_Summary", "Dependency", "Scorecard", "Delivery_Trend",
                 "Exceptions", "Scenario", "Austin_Vendors", "Austin_Categories", "Austin_Prices", "Action_List",
                 "Data_Dictionary", "Validation"]:
        assert name in wb.sheetnames


def test_metrics_are_live_formulas_not_pasted_values(workbook_path):
    wb = load_workbook(workbook_path)
    sc = wb["Scorecard"]
    for col in "BCDEFGHIJKLMNOPQR":
        assert str(sc[f"{col}5"].value).startswith("="), col
    text = " ".join(str(c.value) for row in wb["Supplier_Summary"].iter_rows(min_row=5, max_row=5) for c in row)
    for fn in ["COUNTIFS(", "SUMIFS(", "AVERAGEIFS("]:
        assert fn in text
    assert "XLOOKUP(" in sc["B5"].value
    assert "INDEX(" in str(wb["Dependency"]["Q5"].value) or any("INDEX(" in str(c.value) for c in wb["Dependency"][5])
    lines = wb["Syn_Lines"]
    assert "tblLines" in lines.tables


def test_no_error_literals_or_broken_references(workbook_path):
    wb = load_workbook(workbook_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert "#REF!" not in c.value and "#NAME?" not in c.value, (ws.title, c.coordinate)
                    for sheet in re.findall(r"([A-Za-z_]+)!", c.value):
                        assert sheet in wb.sheetnames, (ws.title, c.coordinate, sheet)


@pytest.mark.slow
def test_recalculated_workbook_matches_python(workbook_path):
    """Recalculate every formula with the `formulas` engine (independent of Excel) and require
    every row of the Validation sheet to MATCH the Python pipeline."""
    formulas = pytest.importorskip("formulas")
    model = formulas.ExcelModel().loads(str(workbook_path)).finish()
    sol = model.calculate()
    vals = {}
    for k, v in sol.items():
        m = re.match(r"'\[.*\](.+)'!([A-Z]+\d+)$", k)
        if m:
            vals[(m.group(1).upper(), m.group(2))] = v.value[0][0]
    results = []
    r = 5
    while ("VALIDATION", f"A{r}") in vals:
        results.append((vals[("VALIDATION", f"A{r}")], vals[("VALIDATION", f"E{r}")], vals[("VALIDATION", f"B{r}")], vals[("VALIDATION", f"C{r}")]))
        r += 1
    assert len(results) >= 50
    mismatches = [x for x in results if x[1] != "MATCH"]
    assert not mismatches, mismatches
